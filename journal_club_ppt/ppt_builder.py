"""Convert presentation plans into PowerPoint files."""
from __future__ import annotations

from dataclasses import dataclass
import random
from typing import Optional

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from .content import ContentSlideData, PresentationPlan, TitleSlideData


@dataclass(frozen=True)
class ColorScheme:
    """Visual identity for a presentation."""

    name: str
    title_background: RGBColor
    content_background: RGBColor
    accent: RGBColor
    title_font: RGBColor
    body_font: RGBColor


COLOR_SCHEMES = [
    ColorScheme(
        name="deep_ocean",
        title_background=RGBColor(0x12, 0x2b, 0x5b),
        content_background=RGBColor(0xf4, 0xf7, 0xfa),
        accent=RGBColor(0x17, 0x6d, 0x81),
        title_font=RGBColor(0xff, 0xff, 0xff),
        body_font=RGBColor(0x23, 0x2f, 0x3d),
    ),
    ColorScheme(
        name="sunset_copper",
        title_background=RGBColor(0x4b, 0x1d, 0x3f),
        content_background=RGBColor(0xff, 0xfa, 0xf2),
        accent=RGBColor(0xd9, 0x6c, 0x06),
        title_font=RGBColor(0xff, 0xff, 0xff),
        body_font=RGBColor(0x3d, 0x2a, 0x26),
    ),
    ColorScheme(
        name="sage_green",
        title_background=RGBColor(0x14, 0x3d, 0x32),
        content_background=RGBColor(0xf3, 0xf6, 0xf3),
        accent=RGBColor(0x6c, 0x99, 0x5b),
        title_font=RGBColor(0xf8, 0xfd, 0xf9),
        body_font=RGBColor(0x24, 0x33, 0x2a),
    ),
    ColorScheme(
        name="storm_gray",
        title_background=RGBColor(0x1f, 0x1f, 0x2e),
        content_background=RGBColor(0xf5, 0xf5, 0xfa),
        accent=RGBColor(0x88, 0x5a, 0xcf),
        title_font=RGBColor(0xff, 0xff, 0xff),
        body_font=RGBColor(0x2f, 0x2f, 0x3a),
    ),
]


def create_presentation(
    plan: PresentationPlan,
    output_path: str,
    author: str,
    seed: Optional[int] = None,
    style_name: Optional[str] = None,
) -> str:
    """Generate a PowerPoint file for the provided plan."""

    prs = Presentation()
    rng = random.Random(seed or hash(plan.title_slide.title) % 2**32)
    color_scheme = select_scheme(style_name, rng)

    prs.core_properties.title = plan.title_slide.title
    prs.core_properties.author = author

    add_title_slide(prs, plan.title_slide, color_scheme)

    for slide_data in plan.slides:
        add_content_slide(prs, slide_data, color_scheme, rng)

    prs.save(output_path)
    return output_path


def select_scheme(style_name: Optional[str], rng: random.Random) -> ColorScheme:
    """Choose a color scheme, optionally honoring a requested name."""

    if style_name:
        for scheme in COLOR_SCHEMES:
            if scheme.name == style_name:
                return scheme
    return rng.choice(COLOR_SCHEMES)


def add_title_slide(prs: Presentation, data: TitleSlideData, scheme: ColorScheme) -> None:
    """Add a styled title slide to the presentation."""

    slide = prs.slides.add_slide(prs.slide_layouts[0])
    set_slide_background(slide, scheme.title_background)
    title_shape = slide.shapes.title
    subtitle_shape = slide.placeholders[1]

    title_tf = title_shape.text_frame
    title_tf.text = data.title
    title_para = title_tf.paragraphs[0]
    title_para.font.size = Pt(44)
    title_para.font.bold = True
    title_para.font.color.rgb = scheme.title_font
    title_para.alignment = PP_ALIGN.LEFT

    subtitle_tf = subtitle_shape.text_frame
    subtitle_tf.clear()
    if data.subtitle_lines:
        first_line, *remaining = data.subtitle_lines
        subtitle_tf.text = first_line
        for line in remaining:
            paragraph = subtitle_tf.add_paragraph()
            paragraph.text = line
    else:
        subtitle_tf.text = "Pharmacy Journal Club"
    for paragraph in subtitle_tf.paragraphs:
        paragraph.font.size = Pt(24)
        paragraph.font.color.rgb = scheme.title_font
        paragraph.alignment = PP_ALIGN.LEFT

    if data.footer_lines:
        footer_box = slide.shapes.add_textbox(
            Inches(0.6),
            prs.slide_height - Inches(1.0),
            prs.slide_width - Inches(1.2),
            Inches(0.6),
        )
        footer_tf = footer_box.text_frame
        footer_tf.clear()
        paragraph = footer_tf.paragraphs[0]
        paragraph.text = " | ".join(data.footer_lines)
        paragraph.font.size = Pt(12)
        paragraph.font.color.rgb = scheme.title_font
        paragraph.alignment = PP_ALIGN.CENTER
        footer_box.fill.background()
        footer_box.line.fill.background()


def add_content_slide(
    prs: Presentation,
    data: ContentSlideData,
    scheme: ColorScheme,
    rng: random.Random,
) -> None:
    """Create a slide with title and bullet content."""

    slide = prs.slides.add_slide(prs.slide_layouts[1])
    set_slide_background(slide, scheme.content_background)

    title_shape = slide.shapes.title
    content_placeholder = slide.placeholders[1]

    title_tf = title_shape.text_frame
    title_tf.text = data.title
    title_para = title_tf.paragraphs[0]
    title_para.font.size = Pt(32)
    title_para.font.bold = True
    title_para.font.color.rgb = scheme.body_font

    text_frame = content_placeholder.text_frame
    text_frame.clear()
    text_frame.margin_left = Inches(0.2)
    text_frame.margin_right = Inches(0.2)
    text_frame.vertical_anchor = MSO_ANCHOR.TOP

    for idx, bullet in enumerate(data.bullets):
        if idx == 0:
            paragraph = text_frame.paragraphs[0]
        else:
            paragraph = text_frame.add_paragraph()
        paragraph.text = bullet
        paragraph.font.size = Pt(20)
        paragraph.font.color.rgb = scheme.body_font
        paragraph.level = 0
        paragraph.line_spacing = 1.2

    if rng.random() < 0.5:
        add_accent_bar(slide, scheme, prs.slide_width, height_inches=0.25 + rng.random() * 0.2)
    elif rng.random() < 0.3:
        add_accent_sidebar(slide, scheme, prs.slide_height, width_inches=0.3 + rng.random() * 0.2)

    if data.notes:
        notes_frame = slide.notes_slide.notes_text_frame
        notes_frame.text = data.notes
    else:
        notes_frame = slide.notes_slide.notes_text_frame
        notes_frame.text = "Review source article for supporting data."


def set_slide_background(slide, color: RGBColor) -> None:
    """Set a solid fill background for a slide."""

    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_accent_bar(slide, scheme: ColorScheme, width, height_inches: float) -> None:
    """Add a horizontal accent bar to a slide."""

    height = Inches(height_inches)
    shape = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE,
        left=0,
        top=0,
        width=width,
        height=height,
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = scheme.accent
    shape.line.fill.background()


def add_accent_sidebar(slide, scheme: ColorScheme, height, width_inches: float) -> None:
    """Add a vertical accent sidebar to a slide."""

    width = Inches(width_inches)
    shape = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE,
        left=0,
        top=0,
        width=width,
        height=height,
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = scheme.accent
    shape.line.fill.background()
